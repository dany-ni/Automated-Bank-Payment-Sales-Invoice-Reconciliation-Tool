import pandas as pd

import itertools
import pandas as pd
from decimal import Decimal
from translations import translate_zuordnungsart, translate_grund



def _is_decimal_series(series):
    return series.apply(lambda v: isinstance(v, Decimal)).any()

# Only change decimal colums to "%.2f"
def _apply_number_format(file_path, columns, decimal_cols, number_format):
    from openpyxl import load_workbook

    wb = load_workbook(file_path)
    ws = wb.active

    col_letters = {
        col_name: idx + 1
        for idx, col_name in enumerate(columns)
        if col_name in decimal_cols
    }

    for col_idx in col_letters.values():
        for row in range(2, ws.max_row + 1):  # row 1 = header
            cell = ws.cell(row=row, column=col_idx)
            if isinstance(cell.value, (int, float)):
                cell.number_format = number_format

    wb.save(file_path)

def save_to_excel(df, file_path, decimal_format="0.00"):
    df = df.copy()

    decimal_cols = [col for col in df.columns if _is_decimal_series(df[col])]
    for col in decimal_cols:
        df[col] = df[col].apply(lambda v: float(v) if isinstance(v, Decimal) else v)

    df.to_excel(file_path, index=False)

    if decimal_cols:
        _apply_number_format(file_path, df.columns, decimal_cols, decimal_format)




def build_type3to11_report(result_df, invoices_df, transactions_df, language="en"):
    rows = []
    for _, g in result_df.iterrows():
        txn_ids, inv_nos = g["transaction_ids"], g["invoice_numbers"]
        if g["shape_label"] != "N:M":
            for txn_id, inv_no in itertools.product(txn_ids, inv_nos):
                rows.append(_flat_row(g, txn_id, inv_no, invoices_df, transactions_df, language))
        else:
            for inv_no in inv_nos:
                rows.append(_flat_row_nm(g, txn_ids, inv_no, invoices_df, transactions_df, language))

    cols = ["rechnungsnummer", "kunde", "auftraggeber_empfaenger", "rechnungsdatum",
            "rechnungsbetrag", "wertstellung", "kontoauszugsbetrag", "betrag_diff",
            "verwendungszweck", "skonto_prozent", "skonto_datum", "zuordnungsart", "grund"]

    report_matching_results_df = pd.DataFrame(rows, columns=cols)

    def _to_date_(v):
        if pd.isna(v):
            return v
        if isinstance(v, str) and ";" in v:
            return v  
        dt = pd.to_datetime(v, errors="coerce")
        return dt.date() if pd.notna(dt) else v

    for date_col in ["rechnungsdatum", "wertstellung", "skonto_datum"]:
        report_matching_results_df[date_col] = report_matching_results_df[date_col].apply(_to_date_)

    return report_matching_results_df


def _flat_row(g, txn_id, inv_no, invoices_df, transactions_df, language):
    inv = invoices_df.loc[inv_no]
    txn = transactions_df.loc[txn_id]
    return {
        "rechnungsnummer": inv_no,
        "kunde": inv["kunde"],
        "auftraggeber_empfaenger": txn["auftraggeber_empfaenger"],
        "rechnungsdatum": inv["rechnungsdatum"],
        "rechnungsbetrag": inv["rechnungsbetrag"],
        "wertstellung": txn["wertstellung"],
        "kontoauszugsbetrag": txn["betrag"],
        "betrag_diff": g["betrag_diff"],   
        "verwendungszweck": txn["verwendungszweck"],
        "skonto_prozent": inv["skonto_prozent"],
        "skonto_datum": inv["skonto_datum"],
        "zuordnungsart": translate_zuordnungsart(g["zuordnungsart"], language),
        "grund": translate_grund(g["grund"], language),
    }


def _flat_row_nm(g, txn_ids, inv_no, invoices_df, transactions_df, language):
    inv = invoices_df.loc[inv_no]
    txns = transactions_df.loc[txn_ids]
    return {
        "rechnungsnummer": inv_no,
        "kunde": inv["kunde"],
        "auftraggeber_empfaenger": "; ".join(txns["auftraggeber_empfaenger"].astype(str)),
        "rechnungsdatum": inv["rechnungsdatum"],
        "rechnungsbetrag": inv["rechnungsbetrag"],
        "wertstellung": "; ".join(txns["wertstellung"].astype(str)),
        "kontoauszugsbetrag": None,   
        "betrag_diff": g["betrag_diff"],
        "verwendungszweck": "; ".join(txns["verwendungszweck"].astype(str)),
        "skonto_prozent": inv["skonto_prozent"],
        "skonto_datum": inv["skonto_datum"],
        "zuordnungsart": translate_zuordnungsart(g["zuordnungsart"], language),
        "grund": translate_grund(g["grund"], language),
    }